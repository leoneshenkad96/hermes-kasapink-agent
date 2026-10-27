"""Durable, approval-aware job execution for Hermes."""

from __future__ import annotations

import csv
import json
import os
import re
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

import daily_report
import engineering_tasks
import local_runner
import monitor
import devops_adapters
import external_adapters
from file_lock import exclusive


STATE_FILE = Path(os.getenv("HERMES_JOB_STATE_FILE", "./data/jobs.json"))
MAX_RETRIES = 2


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _read() -> dict[str, Any]:
    try:
        return json.loads(STATE_FILE.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return {"jobs": {}}


def _write(data: dict[str, Any]) -> None:
    STATE_FILE.parent.mkdir(parents=True, exist_ok=True)
    tmp = STATE_FILE.with_suffix(".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(STATE_FILE)


def _audit(job_id: str, event: str, **details: Any) -> None:
    data = _read()
    record = data.setdefault("jobs", {}).get(job_id)
    if record:
        record.setdefault("events", []).append({"at": _now(), "event": event, **details})
        with exclusive(STATE_FILE):
            _write(data)


def _infer(prompt: str) -> str:
    value = prompt.lower()
    if any(x in value for x in ("external-status", "external status", "github vercel neon", "cek koneksi github", "cek koneksi vercel", "cek koneksi neon")):
        return "external_status"
    if any(x in value for x in ("devops-status", "devops status", "kesiapan devops", "tool adapter", "adapter devops", "cek tool lokal")):
        return "devops_status"
    if any(x in value for x in ("incident diagnosis", "diagnose incident", "diagnosa incident", "insiden", "postmortem")):
        return "incident_plan"
    if any(x in value for x in ("programmer workflow", "debug python", "python debugging", "debugging code")):
        return "programmer_plan"
    if any(x in value for x in ("devops workflow", "ci/cd", "ci cd", "pipeline", "release plan")):
        return "devops_plan"
    if any(x in value for x in ("excel", ".xlsx", ".xls", "spreadsheet", "csv")):
        return "excel_summary"
    if any(x in value for x in ("saham", "stock", "ihsg")):
        return "stock_report"
    if any(x in value for x in ("stok", "erp", "kasapink", "henka")):
        return "erp_report"
    if any(x in value for x in ("mikrotik", "router", "server", "jaringan", "network")):
        return "network_report"
    if any(x in value for x in ("maintenance", "maintenance script", "script")):
        return "maintenance"
    if any(x in value for x in ("codex", "engineering", "task engineering")):
        return "engineering_task"
    if any(x in value for x in ("buka chatgpt", "open chatgpt", "aplikasi")):
        return "app_workflow"
    if any(x in value for x in ("restart", "recovery", "pulihkan", "service mati")):
        return "recovery"
    return "general"


def _plan(kind: str, prompt: str) -> tuple[list[str], str, bool]:
    plans = {
        "erp_report": (["Ambil snapshot ERP read-only", "Rangkum status dan perubahan", "Simpan hasil job"], "Tidak mengubah ERP; hanya membaca data.", False),
        "network_report": (["Cek koneksi dan adapter", "Ambil snapshot Core/LAN/MikroTik read-only", "Rangkum kesehatan jaringan"], "Tidak mengubah router atau konfigurasi.", False),
        "stock_report": (["Ambil data saham watchlist", "Ambil data IHSG", "Buat ringkasan laporan"], "Data eksternal dapat gagal atau tertunda; tidak ada transaksi saham.", False),
        "excel_summary": (["Validasi path file", "Baca tabel secara read-only", "Buat ringkasan dan simpan hasil"], "File tidak diubah; file XLSX memerlukan library spreadsheet terpasang.", False),
        "engineering_task": (["Buat task engineering", "Simpan evidence dan deskripsi", "Tampilkan ID untuk handoff Codex"], "Tidak otomatis merge, deploy, atau mengubah repository.", False),
        "app_workflow": (["Validasi aplikasi allowlist", "Minta approval bila membuka aplikasi", "Jalankan workflow yang tersedia"], "Workflow aplikasi belum arbitrary; hanya aplikasi yang di-allowlist.", True),
        "recovery": (["Cek status Hermes", "Restart Hermes melalui Scheduled Task", "Verifikasi status setelah restart"], "Restart service lokal; tidak mengubah konfigurasi.", True),
        "maintenance": (["Validasi nama script allowlist", "Minta approval", "Jalankan script dengan timeout"], "Script arbitrary diblokir sampai ditambahkan ke allowlist.", True),
        "devops_status": (["Deteksi tool lokal", "Cek versi dan koneksi read-only", "Laporkan adapter GitHub/Vercel/Neon tanpa credential"], "Tidak menjalankan deploy, query write, atau perubahan konfigurasi.", False),
        "external_status": (["Validasi target GitHub", "Baca metadata project Vercel", "Baca metadata project/branch Neon", "Simpan evidence read-only"], "Tidak menjalankan merge, deploy, migration, database write, atau perubahan resource.", False),
        "incident_plan": (["Tetapkan dampak dan scope", "Cek health/process/port/log/dependency", "Susun hipotesis dan recovery", "Tentukan verifikasi"], "Recovery dan perubahan infrastruktur memerlukan approval terpisah.", True),
        "programmer_plan": (["Kumpulkan reproduksi dan traceback", "Periksa diff/dependency/config", "Susun patch terkecil", "Rencanakan regression test"], "Tidak mengubah repository atau menghapus perubahan user.", False),
        "devops_plan": (["Validasi scope dan target", "Build/test/scan", "Stage preview", "Observasi health", "Promote atau rollback setelah approval"], "Deploy, merge, migration, secret, dan rollback production memerlukan approval.", True),
        "general": (["Klasifikasikan permintaan", "Tentukan tool yang diperlukan", "Minta detail atau approval bila belum aman"], "Belum ada handler khusus untuk permintaan ini.", True),
    }
    return plans[kind]


def create(prompt: str, requester_id: int | str) -> dict[str, Any]:
    kind = _infer(prompt)
    steps, risk, approval_required = _plan(kind, prompt)
    job_id = "JOB-" + uuid.uuid4().hex[:8].upper()
    status = "AWAITING_APPROVAL" if approval_required else "PLANNED"
    record = {"id": job_id, "kind": kind, "prompt": prompt[:2000], "requester_id": str(requester_id), "status": status, "steps": steps, "risk": risk, "approval_required": approval_required, "retries": 0, "created_at": _now(), "events": []}
    with exclusive(STATE_FILE):
        data = _read(); data.setdefault("jobs", {})[job_id] = record; _write(data)
    _audit(job_id, "created", kind=kind, approval_required=approval_required)
    return record


def get(job_id: str, requester_id: int | str | None = None) -> dict[str, Any]:
    record = _read().get("jobs", {}).get(job_id.upper())
    if not record:
        raise ValueError("job tidak ditemukan")
    if requester_id is not None and str(record.get("requester_id")) != str(requester_id):
        raise PermissionError("job hanya bisa dilihat oleh pembuatnya")
    return record


def list_jobs(requester_id: int | str, limit: int = 10) -> list[dict[str, Any]]:
    jobs = [x for x in _read().get("jobs", {}).values() if str(x.get("requester_id")) == str(requester_id)]
    return sorted(jobs, key=lambda x: x.get("created_at", ""), reverse=True)[:limit]


def schedule(prompt: str, requester_id: int | str, time_local: str) -> dict[str, Any]:
    if not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", time_local.strip()):
        raise ValueError("waktu harus format HH:MM, misalnya 07:30")
    schedule_id = "SCH-" + uuid.uuid4().hex[:8].upper()
    record = {"id": schedule_id, "prompt": prompt[:2000], "requester_id": str(requester_id), "time_local": time_local.strip(), "enabled": True, "last_run_date": None, "created_at": _now()}
    with exclusive(STATE_FILE):
        data = _read(); data.setdefault("schedules", {})[schedule_id] = record; _write(data)
    return record


def list_schedules(requester_id: int | str) -> list[dict[str, Any]]:
    return [x for x in _read().get("schedules", {}).values() if str(x.get("requester_id")) == str(requester_id) and x.get("enabled")]


def run_due_schedules(now_local: datetime) -> list[dict[str, Any]]:
    due: list[dict[str, Any]] = []
    today = now_local.date().isoformat()
    current = now_local.strftime("%H:%M")
    data = _read()
    for record in data.get("schedules", {}).values():
        if record.get("enabled") and record.get("time_local") == current and record.get("last_run_date") != today:
            job = create(record["prompt"], record["requester_id"])
            result = job
            if not job.get("approval_required"):
                result = run(job["id"], record["requester_id"])
            record["last_run_date"] = today
            due.append(result)
    if due:
        with exclusive(STATE_FILE):
            current_data = _read()
            for record in data.get("schedules", {}).values():
                if record.get("id") in current_data.get("schedules", {}):
                    current_data["schedules"][record["id"]] = record
            _write(current_data)
    return due


def approve(job_id: str, requester_id: int | str) -> dict[str, Any]:
    with exclusive(STATE_FILE):
        data = _read(); record = data.get("jobs", {}).get(job_id.upper())
        if not record: raise ValueError("job tidak ditemukan")
        if str(record.get("requester_id")) != str(requester_id): raise PermissionError("hanya pembuat job yang boleh approve")
        if record.get("status") != "AWAITING_APPROVAL": raise ValueError(f"job berstatus {record.get('status')}")
        record["status"] = "APPROVED"; record["approved_at"] = _now(); _write(data)
    _audit(job_id.upper(), "approved")
    return get(job_id)


def _csv_summary(path: Path) -> dict[str, Any]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))
    return {"file": path.name, "format": "CSV", "rows": len(rows), "columns": list(rows[0].keys()) if rows else [], "sample": rows[:3]}


def _xlsx_summary(path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("library openpyxl belum terpasang") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        header = [str(x) if x is not None else "" for x in (rows[0] if rows else [])]
        sheets.append({"name": sheet.title, "rows": max(0, len(rows) - 1), "columns": header[:30], "sample": [list(row[:10]) for row in rows[1:4]]})
    workbook.close()
    return {"file": path.name, "format": "XLSX", "sheets": sheets}


def _execute(kind: str, prompt: str) -> dict[str, Any]:
    if kind == "erp_report":
        snapshot = monitor.collect(False)
        return {"summary": "Laporan ERP dan sistem berhasil dibuat.", "snapshot": snapshot}
    if kind == "network_report":
        return {"summary": "Laporan jaringan berhasil dibuat.", "network": local_runner.diagnose("network"), "systems": monitor.collect(False)}
    if kind == "stock_report":
        results = daily_report.fetch_stock_sample()
        ihsg = daily_report.fetch_ihsg()
        return {"summary": "Laporan saham berhasil dibuat.", "stocks": results, "ihsg": ihsg}
    if kind == "excel_summary":
        match = re.search(r"(?:[A-Za-z]:[\\/][^\n]+?\.(?:csv|xlsx|xls))", prompt, re.I)
        if not match: raise ValueError("sertakan path file CSV/XLSX yang ingin diringkas")
        path = Path(match.group(0).strip(' `"'))
        if path.suffix.lower() == ".csv":
            return {"summary": "Ringkasan CSV berhasil dibuat.", "data": _csv_summary(path)}
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            return {"summary": "Ringkasan Excel berhasil dibuat secara read-only.", "data": _xlsx_summary(path)}
        raise RuntimeError("format spreadsheet belum didukung; gunakan CSV, XLSX, atau XLSM")
    if kind == "engineering_task":
        title = prompt[:100].strip() or "Hermes engineering job"
        return {"summary": "Task engineering dibuat.", "task": engineering_tasks.create(title, prompt)}
    if kind == "devops_status":
        return {"summary": "Inventaris readiness DevOps read-only berhasil dibuat.", "data": devops_adapters.diagnose()}
    if kind == "external_status":
        return {"summary": "Status external adapter read-only berhasil dibuat.", "data": external_adapters.diagnose()}
    if kind == "incident_plan":
        return {"summary": "Rencana incident diagnosis dibuat.", "plan": devops_adapters.workflow_plan("incident", prompt)}
    if kind == "programmer_plan":
        return {"summary": "Rencana programmer workflow dibuat.", "plan": devops_adapters.workflow_plan("programmer", prompt)}
    if kind == "devops_plan":
        return {"summary": "Rencana DevOps workflow dibuat.", "plan": devops_adapters.workflow_plan("devops", prompt)}
    if kind == "recovery":
        return local_runner.restart_hermes()
    if kind == "maintenance":
        raise RuntimeError("maintenance script belum ada di allowlist; tambahkan script spesifik setelah review")
    if kind == "app_workflow":
        raise RuntimeError("workflow aplikasi belum ada di allowlist; gunakan task ChatGPT startup atau tambahkan workflow spesifik")
    raise RuntimeError("job belum memiliki handler khusus")


def run(job_id: str, requester_id: int | str) -> dict[str, Any]:
    record = get(job_id, requester_id)
    if record.get("approval_required") and record.get("status") != "APPROVED":
        raise PermissionError("job memerlukan approval sebelum dijalankan")
    if record.get("status") in {"SUCCEEDED", "RUNNING"}:
        return record
    with exclusive(STATE_FILE):
        data = _read(); record = data["jobs"][job_id.upper()]; record["status"] = "RUNNING"; record["started_at"] = _now(); _write(data)
    _audit(job_id.upper(), "started")
    last_error = None
    for attempt in range(MAX_RETRIES + 1):
        try:
            result = _execute(record["kind"], record["prompt"])
            with exclusive(STATE_FILE):
                data = _read(); record = data["jobs"][job_id.upper()]; record.update({"status": "SUCCEEDED", "result": result, "finished_at": _now(), "retries": attempt}); _write(data)
            _audit(job_id.upper(), "succeeded", retries=attempt)
            return record
        except Exception as exc:
            last_error = type(exc).__name__
            if attempt < MAX_RETRIES:
                time.sleep(1)
                continue
    with exclusive(STATE_FILE):
        data = _read(); record = data["jobs"][job_id.upper()]; record.update({"status": "FAILED", "error": last_error, "finished_at": _now(), "retries": MAX_RETRIES}); _write(data)
    _audit(job_id.upper(), "failed", error=last_error)
    return get(job_id)
